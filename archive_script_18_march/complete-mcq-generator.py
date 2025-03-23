import os
import sys
import docx
import re
import anthropic
import time
import json
from fpdf import FPDF
from tqdm import tqdm
import argparse

def extract_text_from_docx(file_path):
    """Extract text content from a Word document."""
    try:
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            full_text.append(para.text)
        return '\n'.join(full_text)
    except Exception as e:
        print(f"Error reading Word document: {e}")
        sys.exit(1)

def split_document(text, max_tokens=7500, overlap=500):
    """Split document text into chunks of approximately max_tokens each."""
    # Simple token estimation (not exactly accurate but a good approximation)
    words = text.split()
    tokens_per_word = 1.3  # Rough estimate
    
    chunks = []
    current_chunk = []
    current_token_count = 0
    
    for word in words:
        word_token_estimate = len(word) / 4 + 0.5  # Simple approximation
        if current_token_count + word_token_estimate > max_tokens and current_chunk:
            # Save current chunk
            chunks.append(' '.join(current_chunk))
            
            # Keep some overlap for context
            overlap_words = min(len(current_chunk), int(overlap / tokens_per_word))
            current_chunk = current_chunk[-overlap_words:]
            current_token_count = sum(len(w) / 4 + 0.5 for w in current_chunk)
        
        current_chunk.append(word)
        current_token_count += word_token_estimate
    
    # Don't forget to add the last chunk
    if current_chunk:
        chunks.append(' '.join(current_chunk))
    
    return chunks

def format_chunk_for_claude(chunk, chunk_index, total_chunks, prev_questions=None):
    """Format a document chunk for Claude API."""
    base_message = f"<documents>\n<document index=\"{chunk_index+1}\">\n<source>Transcript ({chunk_index+1}).docx</source>\n<document_content>{chunk}</document_content>\n</document>\n</documents>"
    
    # If this isn't the first chunk, include previous questions to ensure we don't duplicate
    if chunk_index > 0 and prev_questions:
        base_message += f"\n\nPlease don't duplicate these questions that were already generated from previous parts of the transcript:\n{prev_questions}"
    
    if chunk_index > 0:
        base_message += "\n\nMake sure to combine these new 5 MCQ questions with the previous questions, for a total set of questions."
    
    return base_message

def get_claude_prompt():
    """Return the prompt template for Claude."""
    return """This is a Transcript of a course, generate set of 5 mcq questions for a course end quiz from this Material.  
1. Make sure not to pick up any questions where it's not part of the course or related to course topic.  
2. Make sure to properly distribute the questions from full script 
3. Make sure to revalidate each questions and answers again  
4. The questions should be direct and clear, question answer keywords should be present in the transcript  
5. Each questions should also mention the timestamp range from where it is picked up 
6. Questions should be only related to course content; remove any miscellaneous questions.  
7. Please also give explanation of each questions answer 
8. Also calculate input & output token at the end
9. Use markdown formatting for the questions
10. Each question should be numbered properly (Question 1, Question 2, etc.)
11. Make sure to include 4 answer choices (A, B, C, D) for each question
12. Mark the correct answer by stating "Answer: [correct option]" after the choices
13. Add an explanation of why the answer is correct after stating the answer"""

def call_claude_api(api_key, content, model="claude-3-5-sonnet-20240620"):
    """Call Claude API to generate MCQ questions."""
    client = anthropic.Anthropic(api_key=api_key)
    
    try:
        message = client.messages.create(
            model=model,
            max_tokens=4096,
            temperature=0.2,
            system="You are a helpful AI assistant that creates high-quality multiple choice questions based on course transcripts.",
            messages=[
                {"role": "user", "content": f"{get_claude_prompt()}\n\n{content}"}
            ]
        )
        return message.content[0].text
    except Exception as e:
        print(f"Error calling Claude API: {e}")
        return None

def extract_mcq_from_response(response):
    """Extract only the MCQ questions from Claude's response."""
    # Try to locate MCQ questions with regex patterns
    # Look for patterns like "Question 1", "## Question 1", numbered lists, etc.
    
    # First, let's check if there's a markdown-style section with questions
    mcq_section = ""
    
    # Try to extract markdown-formatted questions
    if "## Question" in response or "# Question" in response:
        # It's probably in markdown format
        questions_match = re.search(r'(#+\s*Question\s*\d[\s\S]*?)(?=\n*Input tokens|\n*$)', response, re.IGNORECASE)
        if questions_match:
            mcq_section = questions_match.group(1)
    
    # If that didn't work, try to find numbered questions (1., 2., etc.)
    if not mcq_section:
        questions_match = re.search(r'((?:^\s*\d+\.[\s\S]*?){3,})', response, re.MULTILINE)
        if questions_match:
            mcq_section = questions_match.group(1)
    
    # If we still don't have MCQs, just use the whole response
    if not mcq_section:
        mcq_section = response
    
    return mcq_section

def combine_mcq_responses(responses):
    """Combine multiple MCQ responses into a single comprehensive set."""
    combined = ""
    question_count = 0
    
    for response in responses:
        # Extract questions and reformat numbering if needed
        clean_response = extract_mcq_from_response(response)
        
        if not clean_response:
            continue
            
        # Change question numbers to continue from previous chunks
        # Replace "Question X" or "## Question X" with proper numbering
        lines = clean_response.split('\n')
        new_lines = []
        
        for line in lines:
            # Match "Question X" or "## Question X" patterns
            question_match = re.search(r'(#+\s*)?Question\s+(\d+)', line, re.IGNORECASE)
            if question_match:
                prefix = question_match.group(1) or ''
                question_count += 1
                new_line = re.sub(r'(#+\s*)?Question\s+\d+', f'{prefix}Question {question_count}', line)
                new_lines.append(new_line)
            else:
                new_lines.append(line)
        
        # Add to combined response
        if combined:
            combined += "\n\n" + '\n'.join(new_lines)
        else:
            combined = '\n'.join(new_lines)
    
    # Add a title
    combined = "# Multiple Choice Questions from Course Transcript\n\n" + combined
    
    return combined

def create_pdf(mcq_content, output_file="mcq_questions.pdf"):
    """Create a PDF with all the MCQ questions using standard fonts."""
    # Initialize PDF
    pdf = FPDF()
    pdf.add_page()
    # Use standard fonts that come with FPDF
    pdf.set_font('Arial', '', 11)
    
    # Set title
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, "Multiple Choice Questions from Course Transcript", 0, 1, 'C')
    pdf.ln(5)
    
    # Process content
    lines = mcq_content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Skip empty lines
        if not line:
            pdf.ln(5)
            i += 1
            continue
        
        # Question headers (## Question X or # Question X or Question X)
        question_match = re.search(r'(#+\s*)?Question\s+(\d+)', line, re.IGNORECASE)
        if question_match:
            # Add some spacing before questions (except the first one)
            if question_match.group(2) != "1":
                pdf.ln(10)
                
            pdf.set_font('Arial', 'B', 14)
            # Clean the line from markdown
            clean_line = re.sub(r'^#+\s*', '', line)
            pdf.multi_cell(0, 7, clean_line)
            pdf.set_font('Arial', '', 11)
            i += 1
            continue
        
        # Timestamp
        timestamp_match = re.search(r'\[Timestamp:', line)
        if timestamp_match:
            pdf.set_font('Arial', 'I', 10)
            pdf.multi_cell(0, 5, line)
            pdf.set_font('Arial', '', 11)
            i += 1
            continue
        
        # Answer choices (A), B), C), D) or A. B. C. D.)
        choice_match = re.search(r'^([A-D][\)\.])', line)
        if choice_match:
            pdf.multi_cell(0, 6, line)
            i += 1
            continue
        
        # Answer line
        if line.startswith("Answer:") or line.startswith("**Answer:"):
            pdf.set_font('Arial', 'B', 11)
            clean_line = re.sub(r'\*\*', '', line)  # Remove markdown bold
            pdf.multi_cell(0, 6, clean_line)
            pdf.set_font('Arial', '', 11)
            i += 1
            continue
        
        # Explanation
        if line.startswith("Explanation:") or line.startswith("**Explanation:"):
            pdf.set_font('Arial', 'I', 11)
            pdf.multi_cell(0, 6, "Explanation:")
            pdf.set_font('Arial', '', 11)
            
            # Gather the explanation text (which might span multiple lines)
            explanation_text = ""
            i += 1
            while i < len(lines) and not (lines[i].strip().startswith("## Question") or lines[i].strip().startswith("# Question") or lines[i].strip().startswith("Question")):
                if lines[i].strip():  # Skip empty lines in explanation
                    explanation_text += lines[i].strip() + " "
                i += 1
                
                # Break if we've reached the end or another question/section
                if i >= len(lines) or re.search(r'(#+\s*)?Question\s+(\d+)', lines[i].strip(), re.IGNORECASE):
                    break
            
            # Clean explanation text from markdown
            explanation_text = re.sub(r'\*\*', '', explanation_text)  # Remove bold
            explanation_text = re.sub(r'_([^_]+)_', r'\1', explanation_text)  # Remove italic
            
            pdf.multi_cell(0, 6, explanation_text)
            continue
        
        # Regular text
        pdf.multi_cell(0, 6, line)
        i += 1
    
    # Save PDF
    try:
        pdf.output(output_file)
        print(f"PDF saved as {output_file}")
        return True
    except Exception as e:
        print(f"Error creating PDF: {e}")
        # Fallback to simple text file
        print("Creating a text file instead...")
        with open(output_file.replace('.pdf', '.txt'), 'w', encoding='utf-8') as f:
            f.write(mcq_content)
        return False

def convert_mcq_to_json(mcq_content):
    """Convert markdown MCQ content to a JSON structure."""
    questions = []
    current_question = None
    current_choices = []
    current_section = None  # To track what section we're parsing
    
    lines = mcq_content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i].strip()
        
        # Skip empty lines
        if not line:
            i += 1
            continue
        
        # Question headers - detect both markdown and plain formats
        question_match = re.search(r'(#+\s*)?Question\s+(\d+)', line, re.IGNORECASE)
        
        if question_match:
            # If we already have a question, save it before starting a new one
            if current_question:
                if current_choices:
                    current_question['choices'] = current_choices
                questions.append(current_question)
            
            # Start building new question
            question_number = question_match.group(2)
            question_text = re.sub(r'(#+\s*)?Question\s+\d+[\.:]?\s*', '', line).strip()
            
            # Initialize new question object
            current_question = {
                "id": int(question_number),
                "question": question_text,
                "timestamp": ""
            }
            current_choices = []
            current_section = "question"
            i += 1
            continue
            
        # Timestamp - usually right after question
        timestamp_match = re.search(r'\[Timestamp:([^\]]+)\]', line)
        if timestamp_match and current_question:
            current_question['timestamp'] = line.strip()
            i += 1
            continue
            
        # Answer choices - handle A), A., or A format
        choice_match = re.search(r'^([A-D])[\s\.\)]+(.+)$', line)
        if choice_match and current_question:
            choice_letter = choice_match.group(1)
            choice_text = choice_match.group(2).strip()
            
            # Clean up markdown formatting
            choice_text = re.sub(r'\*\*|\*|_', '', choice_text)
            
            current_choices.append({
                "letter": choice_letter,
                "text": choice_text
            })
            current_section = "choices"
            i += 1
            continue
            
        # Correct answer line
        answer_match = re.search(r'(Answer|Answer:)\s*\**([A-D])\**', line, re.IGNORECASE)
        if answer_match and current_question:
            correct_answer = answer_match.group(2)
            current_question['correctAnswer'] = correct_answer
            current_section = "answer"
            i += 1
            continue
            
        # Explanation starter
        explanation_match = re.search(r'(Explanation|Explanation:)', line, re.IGNORECASE)
        if explanation_match and current_question:
            # Initialize explanation
            explanation_text = ""
            
            # Move past the "Explanation:" line
            i += 1
            current_section = "explanation"
            
            # Collect all lines until next question or end
            while i < len(lines):
                line = lines[i].strip()
                
                # Break if we hit the next question
                if re.search(r'(#+\s*)?Question\s+\d+', line, re.IGNORECASE):
                    break
                    
                # Add non-empty lines to explanation
                if line:
                    explanation_text += line + " "
                    
                i += 1
                
            # Clean up the explanation
            explanation_text = explanation_text.strip()
            explanation_text = re.sub(r'\*\*|\*|_', '', explanation_text)  # Remove markdown
            
            current_question['explanation'] = explanation_text
            continue
            
        # If we're in the explanation section but didn't catch it above
        if current_section == "explanation" and current_question and "explanation" not in current_question:
            explanation_text = line
            current_question['explanation'] = explanation_text
            
        # Move to next line if no patterns matched
        i += 1
    
    # Don't forget to add the last question
    if current_question:
        if current_choices:
            current_question['choices'] = current_choices
        questions.append(current_question)
    
    # Revalidate the questions to ensure they have all required fields
    validated_questions = []
    for q in questions:
        # Only include questions that have the minimal required structure
        if 'question' in q and 'choices' in q and len(q['choices']) > 0:
            # Ensure all questions have an explanation field
            if 'explanation' not in q:
                q['explanation'] = ""
            
            # Ensure all have a timestamp
            if 'timestamp' not in q:
                q['timestamp'] = ""
                
            # Ensure correctAnswer exists
            if 'correctAnswer' not in q:
                # If we can find it in the choices
                for choice in q['choices']:
                    if "correct" in choice['text'].lower() or "answer" in choice['text'].lower():
                        q['correctAnswer'] = choice['letter']
                        break
            
            validated_questions.append(q)
    
    return {"questions": validated_questions}

def main():
    # Set up argument parser
    parser = argparse.ArgumentParser(description='Process a Word document, split it, generate MCQs using Claude API, and compile into a PDF.')
    parser.add_argument('--docx', '-d', required=True, help='Path to the Word document')
    parser.add_argument('--api-key', '-k', required=True, help='Anthropic API key')
    parser.add_argument('--output', '-o', default='course_mcq_questions.pdf', help='Output PDF file name')
    parser.add_argument('--max-chunks', '-c', type=int, default=4, help='Maximum number of chunks to process')
    parser.add_argument('--model', '-m', default='claude-3-5-sonnet-20240620', help='Claude model to use')
    
    args = parser.parse_args()
    
    # Extract text from Word document
    print(f"Reading document: {args.docx}")
    doc_text = extract_text_from_docx(args.docx)
    
    # Split the document
    print("Splitting document into chunks...")
    chunks = split_document(doc_text)
    print(f"Document split into {len(chunks)} chunks")
    
    # Limit chunks if requested
    max_chunks = min(len(chunks), args.max_chunks)
    chunks = chunks[:max_chunks]
    
    # Process each chunk with Claude API
    mcq_responses = []
    previous_questions = ""
    
    print(f"Processing {max_chunks} chunks with Claude API...")
    for i, chunk in enumerate(tqdm(chunks)):
        print(f"\nProcessing chunk {i+1}/{max_chunks}")
        
        # Format chunk for Claude
        formatted_chunk = format_chunk_for_claude(chunk, i, max_chunks, previous_questions)
        
        # Call Claude API
        response = call_claude_api(args.api_key, formatted_chunk, args.model)
        
        if response:
            # Extract only the MCQ portion
            mcq_content = extract_mcq_from_response(response)
            mcq_responses.append(mcq_content)
            
            # Add to previous questions for context in next chunk
            previous_questions += f"\n\n{mcq_content}"
            
            # Sleep to respect API rate limits
            time.sleep(2)
        else:
            print(f"Failed to get response for chunk {i+1}")
    
    # Combine all MCQ responses
    print("Combining all MCQ questions...")
    combined_mcq = combine_mcq_responses(mcq_responses)
    
    # Create PDF with all MCQs
    print("Creating PDF with all MCQ questions...")
    create_pdf(combined_mcq, args.output)
    
    # Convert to JSON and save
    print("Converting to JSON format...")
    json_output = convert_mcq_to_json(combined_mcq)
    json_filename = args.output.replace('.pdf', '.json')
    with open(json_filename, 'w', encoding='utf-8') as f:
        json.dump(json_output, f, indent=2)
    print(f"JSON saved as {json_filename}")
    
    print("Done!")

if __name__ == "__main__":
    main()
